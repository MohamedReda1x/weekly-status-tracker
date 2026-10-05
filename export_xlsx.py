import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

HEAD, CALC, GREEN, ORANGE = PatternFill("solid", fgColor="DDE6F0"), PatternFill("solid", fgColor="E8EDF3"), PatternFill("solid", fgColor="A9D18E"), PatternFill("solid", fgColor="F4B183")
thin = Side(style="thin", color="B8C2CE"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MN = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
ml = lambda k: f"{MN[int(k[5:])-1]}-{k[2:4]}"
num = lambda m: m / 1000

def build(d):
    t, P, A = d["tracker"], d["periods"], d["activities"]; wb = Workbook(); ws = wb.active; ws.title = "Weekly Status"
    ws["A1"] = "Weekly Process Status Report"; ws["A1"].font = Font(bold=True, size=14)
    for r, (k, v) in enumerate([("Employee", d["owner"]["name"]), ("Work package", t["work_package"]), ("Mission start", t["start_date"]), ("Mission end", t["end_date"] or "open"), ("Periods exported", f'{d["range"]["from"]} to {d["range"]["to"]}')], 2):
        ws.cell(r, 1, k).font = Font(bold=True); ws.cell(r, 2, v)
    ws["D2"], ws["D3"] = "Completed", "WIP"; ws["D2"].fill, ws["D3"].fill = GREEN, ORANGE
    h = 9; left = ["Activity details", "Status", "Affected projects", "Deliverables / Functions", "Estimation", "Project Progress in %"]; c0 = len(left) + 1
    for i, n in enumerate(left, 1): ws.cell(h, i, n); ws.merge_cells(start_row=h, start_column=i, end_row=h + 2, end_column=i)
    col, start = {}, 0
    for i, p in enumerate(P):
        col[p["id"]] = c0 + i
        if i == len(P) - 1 or P[i + 1]["month_key"] != p["month_key"]:
            ws.cell(h, c0 + start, ml(p["month_key"]) + ("" if d["month_complete"][p["month_key"]] else "*")); ws.merge_cells(start_row=h, start_column=c0 + start, end_row=h, end_column=c0 + i); start = i + 1
        ws.cell(h + 1, c0 + i, f"WEEK{p['iso_week']}"); ws.cell(h + 2, c0 + i, f"{p['capacity']} day" + ("s" if p["capacity"] > 1 else ""))
    for r in range(h, h + 3):
        for c in range(1, c0 + len(P)): x = ws.cell(r, c); x.fill, x.font, x.border, x.alignment = HEAD, Font(bold=True), BOX, Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = h + 3
    for a in A:
        vals = [a["details"] + (" (archived)" if a["archived"] else ""), a["status"], a["affected"], a["deliverables"], a["estimation"], a["progress"] / 100]
        for i, v in enumerate(vals, 1): ws.cell(r, i, v).border = BOX
        ws.cell(r, 2).fill = GREEN if a["status"] == "Completed" else ORANGE; ws.cell(r, 6).number_format = "0%"
        for p in P:
            x = ws.cell(r, col[p["id"]], num(d["entries"].get(a["id"], {}).get(p["id"], 0)) or None); x.number_format = "0.###"; x.border = BOX
        r += 1
    r += 1
    def line(label, vals, calc, fmt="0.###"):
        nonlocal r
        ws.cell(r, 1, label).font = Font(bold=True); ws.cell(r, 1).alignment = Alignment(horizontal="right"); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c0 - 1)
        for p, v in vals: x = ws.cell(r, col[p["id"]], v); x.number_format = fmt; x.border = BOX; x.fill = CALC if calc else PatternFill()
        r += 1
    line("Leave / Holidays :", [(p, num(d["info"].get(p["id"], {}).get("leave", 0)) or None) for p in P], False)
    line("PH holidays :", [(p, num(d["info"].get(p["id"], {}).get("ph", 0)) or None) for p in P], False)
    line("Actual days spent on activity / week :", [(p, num(d["col_totals"].get(p["id"], 0))) for p in P], True)
    ws.cell(r, 1, "Monthwise (in days) :").font = Font(bold=True); ws.cell(r, 1).alignment = Alignment(horizontal="right"); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c0 - 1)
    i = 0
    while i < len(P):
        j = i
        while j + 1 < len(P) and P[j + 1]["month_key"] == P[i]["month_key"]: j += 1
        x = ws.cell(r, c0 + i, num(d["month_totals"][P[i]["month_key"]])); x.number_format = "0.###"; x.fill = CALC; x.font = Font(bold=True); x.alignment = Alignment(horizontal="center")
        if j > i: ws.merge_cells(start_row=r, start_column=c0 + i, end_row=r, end_column=c0 + j)
        i = j + 1
    if not all(d["month_complete"].values()): ws.cell(r + 1, 1, "* month only partly included in this export: its totals cover the exported periods only.").font = Font(italic=True)
    for i, w in enumerate([44, 12, 18, 26, 16, 12], 1): ws.column_dimensions[L(i)].width = w
    for i in range(len(P)): ws.column_dimensions[L(c0 + i)].width = 9
    ws.freeze_panes = ws.cell(h + 3, 3)
    b = wb.create_sheet("Monthly BL"); ms = sorted(d["month_totals"]); b.append(["Activity", "Status"] + [f"BL {ml(m)}" + ("" if d["month_complete"][m] else "*") for m in ms])
    for a in A: b.append([a["details"] + (" (archived)" if a["archived"] else ""), a["status"]] + [num(d["bl"].get(a["id"], {}).get(m, 0)) or None for m in ms])
    b.append(["Total BL", ""] + [num(sum(d["bl"].get(a["id"], {}).get(m, 0) for a in A)) for m in ms]); b.append(["Monthwise (in days)", ""] + [num(d["month_totals"][m]) for m in ms])
    for row in b.iter_rows():
        for x in row: x.border = BOX; x.number_format = "0.###"
    for x in b[1] + b[b.max_row - 1] + b[b.max_row]: x.font = Font(bold=True); x.fill = HEAD if x.row == 1 else CALC
    b.column_dimensions["A"].width = 44
    s = wb.create_sheet("Summary report"); s.append(["Date", "Author", "Scope", "Comment"])
    for c in d["comments"]: s.append([c["created"].replace("T", " ")[:16], c["author"], ", ".join(f"WEEK{w}/{y}" for y, w in c["scopes"]) or "General comment", c["text"]])
    for x in s[1]: x.font = Font(bold=True); x.fill = HEAD
    for k, w in zip("ABCD", (18, 20, 34, 90)): s.column_dimensions[k].width = w
    for row in s.iter_rows(min_row=2):
        for x in row: x.alignment = Alignment(wrap_text=True, vertical="top")
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
