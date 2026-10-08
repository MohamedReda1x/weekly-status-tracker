"""Security tests with mocked SQL; no database or migrations."""
import os
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import MagicMock, patch

from fastapi import HTTPException, Request
from fastapi.testclient import TestClient
import app


class SecurityTests(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, {"COOKIE_SECURE": "false"})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.connect = patch.object(app.psycopg, "connect",
                                    side_effect=AssertionError("Real database access forbidden"))
        self.connect.start()
        self.addCleanup(self.connect.stop)
        app.FAILS.clear()
        app.IP_ATTEMPTS.clear()
        self.addCleanup(app.FAILS.clear)
        self.addCleanup(app.IP_ATTEMPTS.clear)
        self.sql = MagicMock()
        self.sql.execute.return_value.fetchone.return_value = None
        self.db = patch.object(app, "db")
        self.db.start().return_value.__enter__.return_value = self.sql
        self.addCleanup(self.db.stop)
        self.client = TestClient(app.app)  # No lifespan: init() is not called.
        self.addCleanup(self.client.close)

    def attempt(self, email, **kwargs):
        return self.client.post("/api/login", json={"email": email, "password": "wrong"}, **kwargs)

    def test_ip_limit_across_emails_and_expiry(self):
        with patch.object(app.time, "monotonic", return_value=1000):
            for i in range(app.IP_LOGIN_LIMIT):
                self.assertEqual(self.attempt(f"u{i}@example.org").status_code, 401)
            calls = self.sql.execute.call_count
            blocked = self.attempt("another@example.org")
            self.assertEqual(blocked.status_code, 429)
            self.assertEqual(blocked.headers["Cache-Control"], "no-store")
            self.assertEqual(self.sql.execute.call_count, calls)
        with patch.object(app.time, "monotonic", return_value=1901):
            self.assertEqual(self.attempt("fresh@example.org").status_code, 401)

    def test_email_limit_is_preserved(self):
        for _ in range(5):
            self.assertEqual(self.attempt("same@example.org").status_code, 401)
        self.assertEqual(self.attempt("same@example.org").status_code, 429)

    def test_ips_are_independent_and_header_cannot_bypass(self):
        app.IP_ATTEMPTS["testclient"] = [app.time.monotonic()] * app.IP_LOGIN_LIMIT
        self.assertEqual(self.attempt("a@example.org",
                         headers={"X-Forwarded-For": "198.51.100.1"}).status_code, 429)
        req = Request({"type": "http", "client": ("198.51.100.2", 1234)})
        app.limit_login_ip(req)
        self.assertEqual(len(app.IP_ATTEMPTS["198.51.100.2"]), 1)

    def test_concurrent_ip_limit(self):
        req = Request({"type": "http", "client": ("198.51.100.3", 1234)})
        def attempt(_):
            try:
                app.limit_login_ip(req)
                return 200
            except HTTPException as error:
                return error.status_code
        with ThreadPoolExecutor(max_workers=8) as pool:
            codes = list(pool.map(attempt, range(40)))
        self.assertEqual(codes.count(200), app.IP_LOGIN_LIMIT)
        self.assertEqual(codes.count(429), 40 - app.IP_LOGIN_LIMIT)

    def test_security_headers_and_disabled_docs(self):
        self.assertIsNone(app.app.docs_url)
        self.assertIsNone(app.app.redoc_url)
        self.assertIsNone(app.app.openapi_url)
        for path in ("/docs", "/redoc", "/openapi.json"):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 404)
            self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
            self.assertEqual(response.headers["X-Frame-Options"], "DENY")
            self.assertEqual(response.headers["Referrer-Policy"], "same-origin")
            self.assertNotIn("Strict-Transport-Security", response.headers)
            self.assertNotIn("Cache-Control", response.headers)
        with patch.dict(os.environ, {"COOKIE_SECURE": "true"}):
            response = self.client.get("/api/health")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["Cache-Control"], "no-store")
            self.assertEqual(response.headers["Strict-Transport-Security"], "max-age=31536000")
        with patch.object(app, "db", side_effect=RuntimeError("offline")):
            response = self.client.get("/api/health")
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.headers["Cache-Control"], "no-store")

    def test_secure_startup_validation_precedes_init(self):
        with patch.object(app, "init") as init:
            for url, password in (
                ("http://example.org", "Long-unique-password"),
                ("", "Long-unique-password"),
                ("https://example.org", "ChangeMe-12345"),
                ("https://example.org", "Demo-12345"),
                ("https://example.org", "short"),
            ):
                with self.subTest(url=url, password=password):
                    with patch.dict(os.environ, {"COOKIE_SECURE": "true",
                                    "APP_URL": url, "ADMIN_PASSWORD": password}):
                        with self.assertRaises(RuntimeError):
                            with TestClient(app.app):
                                pass
            init.assert_not_called()
            with patch.dict(os.environ, {"COOKIE_SECURE": "true",
                            "APP_URL": "https://example.org",
                            "ADMIN_PASSWORD": "Long-unique-password"}):
                with TestClient(app.app):
                    pass
            init.assert_called_once()
            init.reset_mock()
            with patch.dict(os.environ, {"COOKIE_SECURE": "false",
                            "APP_URL": "http://localhost:8000", "ADMIN_PASSWORD": ""}):
                app._startup()
            init.assert_called_once()



    def test_new_entries_outside_mission_dates(self):
        for start, end, month, week in (
            ("2026-09-07", "2026-09-11", "2026-09", 37),
            ("2026-12-07", "2026-12-11", "2026-12", 50),
        ):
            with self.subTest(period=start):
                period = dict(id=1, start=start, end=end, month_key=month,
                              iso_year=2026, iso_week=week, capacity=5)
                def execute(sql, params=()):
                    result = MagicMock()
                    if sql.startswith("SELECT id, p_start"):
                        result.__iter__.return_value = iter([period])
                    elif sql.startswith("SELECT * FROM activities"):
                        result.fetchone.return_value = dict(id=7, archived=0, details="Review")
                    elif sql.startswith("SELECT milli"):
                        result.fetchone.return_value = None
                    elif sql.startswith("SELECT COALESCE(SUM"):
                        result.fetchone.return_value = {"v": 1000}
                    return result
                self.sql.execute.side_effect = execute
                with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(
                    app, "get_tracker", return_value={"start_date": "2026-11-01", "end_date": "2026-11-30"}
                ):
                    result = app.save(1, app.Save(changes=[dict(type="cell", activity_id=7,
                                      period_id=1, old=0, new="1")]), Request({"type": "http"}))
                self.assertEqual(result, {"ok": True, "changes": 1})
                statements = [call.args[0] for call in self.sql.execute.call_args_list]
                self.assertIn("COMMIT", statements)
                self.assertTrue(any(sql.startswith("INSERT INTO entries") for sql in statements))


    def test_professional_role_cannot_cross_trades(self):
        def execute(sql, params=()):
            result = MagicMock()
            result.fetchone.return_value = None if sql.startswith("SELECT 1 FROM professional_roles") else {"id": 1}
            return result
        self.sql.execute.side_effect = execute
        for b in (app.Affiliation(trade_id=2, professional_role_id=1),
                  app.Affiliation(professional_role_id=1)):
            with self.assertRaises(HTTPException) as error:
                app.validate_affiliation(self.sql, b)
            self.assertEqual(error.exception.status_code, 422)
        app.validate_affiliation(self.sql, app.Affiliation(trade_id=2, client_id=1))

    def test_employee_cannot_manage_affiliations(self):
        with patch.object(app, "cur_user", return_value={"id": 1, "role": "employee"}):
            for action in (
                lambda: app.affiliation_catalog(Request({"type": "http"})),
                lambda: app.add_affiliation_item(app.CatalogItem(kind="client", name="Other"), Request({"type": "http"})),
                lambda: app.set_affiliation(1, app.Affiliation(), Request({"type": "http"})),
            ):
                with self.assertRaises(HTTPException) as error: action()
                self.assertEqual(error.exception.status_code, 403)
        self.sql.execute.assert_not_called()

    def test_leave_dates_validation(self):
        period = {"start": "2026-10-05", "end": "2026-10-09", "capacity": 5}
        self.assertEqual(app.checked_leave_dates(["2026-10-07", "2026-10-05"], period),
                         ["2026-10-05", "2026-10-07"])
        for dates in (["2026-10-04"], ["2026-10-12"], ["bad"],
                      ["2026-10-05", "2026-10-05"]):
            with self.subTest(dates=dates), self.assertRaises(HTTPException) as error:
                app.checked_leave_dates(dates, period)
            self.assertEqual(error.exception.status_code, 422)

    def test_leave_dates_conflict_rolls_back_without_deleting(self):
        from datetime import date
        def execute(sql, params=()):
            result = MagicMock()
            if sql.startswith("SELECT " + app.PSEL):
                result.fetchone.return_value = {"start": "2026-10-05", "end": "2026-10-09", "capacity": 5}
            if sql.startswith("SELECT leave_date"):
                result.__iter__.return_value = iter([{"leave_date": date(2026, 10, 6)}])
            return result
        self.sql.execute.side_effect = execute
        with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(app, "get_tracker"):
            with self.assertRaises(HTTPException) as error:
                app.write_leave_dates(1, 1, app.LeaveDates(old=[], dates=["2026-10-05"]), Request({"type": "http"}))
        self.assertEqual(error.exception.status_code, 409)
        statements = [call.args[0] for call in self.sql.execute.call_args_list]
        self.assertIn("ROLLBACK", statements)
        self.assertFalse(any(s.startswith("DELETE") or s.startswith("INSERT") for s in statements))

    def test_leave_dates_write_does_not_change_grid_duration(self):
        def execute(sql, params=()):
            result = MagicMock()
            if sql.startswith("SELECT " + app.PSEL):
                result.fetchone.return_value = {"start": "2026-10-05", "end": "2026-10-09", "capacity": 5}
            if sql.startswith("SELECT leave_date"): result.__iter__.return_value = iter([])
            return result
        self.sql.execute.side_effect = execute
        with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(app, "get_tracker") as tracker:
            result = app.write_leave_dates(1, 1, app.LeaveDates(old=[], dates=["2026-10-05"]), Request({"type": "http"}))
        self.assertEqual(result, {"dates": ["2026-10-05"]})
        tracker.assert_called_once_with(self.sql, 1, {"id": 1}, lock=True)
        statements = [call.args[0] for call in self.sql.execute.call_args_list]
        self.assertIn("COMMIT", statements)
        self.assertTrue(any(s.startswith("INSERT INTO history") for s in statements))
        self.assertFalse(any("INSERT INTO info" in s or "UPDATE info" in s or "DELETE FROM info" in s for s in statements))

    def test_current_week_comment_scope_ignores_changed_past_week(self):
        period = dict(id=1, start="2026-09-07", end="2026-09-11", month_key="2026-09",
                      iso_year=2026, iso_week=37, capacity=5)
        def execute(sql, params=()):
            result = MagicMock()
            if sql.startswith("SELECT id, p_start"): result.__iter__.return_value = iter([period])
            elif sql.startswith("SELECT milli"): result.fetchone.return_value = None
            elif sql.startswith("INSERT INTO comments"): result.fetchone.return_value = {"id": 11}
            return result
        self.sql.execute.side_effect = execute
        with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(app, "get_tracker"):
            app.save(1, app.Save(changes=[dict(type="info", kind="leave", period_id=1, old=0, new="1")],
                     comment="Current week", weeks=["2026-41"], comment_scope_explicit=True), Request({"type": "http"}))
        scopes = [call.args[1] for call in self.sql.execute.call_args_list if call.args[0].startswith("INSERT INTO comment_scopes")]
        self.assertEqual(scopes, [(11, 2026, 41)])


    def test_leave_dates_cannot_be_written_by_another_employee(self):
        self.sql.execute.return_value.fetchone.return_value = {"id": 1, "owner_id": 2}
        with patch.object(app, "cur_user", return_value={"id": 1, "role": "employee"}):
            with self.assertRaises(HTTPException) as error:
                app.write_leave_dates(1, 1, app.LeaveDates(old=[], dates=["2026-10-05"]), Request({"type": "http"}))
        self.assertEqual(error.exception.status_code, 404)
        statements = [call.args[0] for call in self.sql.execute.call_args_list]
        self.assertIn("ROLLBACK", statements)
        self.assertFalse(any(s.startswith("DELETE") or s.startswith("INSERT") for s in statements))

    def test_migration_0004_declarations_without_database(self):
        import runpy
        from pathlib import Path
        from alembic import op
        path = Path(app.ROOT) / "migrations" / "versions" / "0004_affiliations_leave_dates.py"
        with patch.object(op, "execute") as execute:
            migration = runpy.run_path(str(path))
            self.assertEqual(migration["down_revision"], "0003")
            migration["upgrade"]()  # execute is mocked: no engine, connection or SQL execution.
        sql = "\n".join(call.args[0] for call in execute.call_args_list)
        self.assertIn("FOREIGN KEY(professional_role_id,trade_id)", sql)
        self.assertIn("CHECK(professional_role_id IS NULL OR trade_id IS NOT NULL)", sql)
        for name in ("Train Control", "Train System", "SwQA", "Sub-SE", "Alstom Petite-Forêt"):
            self.assertIn(name, sql)
        statements = [call.args[0].strip().upper() for call in execute.call_args_list]
        self.assertFalse(any(s.startswith(("DELETE ", "DROP ", "TRUNCATE ")) for s in statements))


    def test_recoverable_deletion_and_restoration_keep_entries_and_history(self):
        activity = dict(id=7, tracker_id=1, details="Populated", deleted=0, archived=0)
        def execute(sql, params=()):
            result = MagicMock()
            if sql.startswith("SELECT * FROM activities"): result.fetchone.return_value = dict(activity)
            elif sql.startswith("UPDATE activities SET deleted"): activity["deleted"] = params[0]
            elif sql.startswith("SELECT p.iso_week"): result.__iter__.return_value = iter([])
            return result
        self.sql.execute.side_effect = execute
        with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(app, "get_tracker") as tracker:
            app.delete_activity(7, Request({"type": "http"}))
            self.assertEqual(activity["deleted"], 1)
            app.delete_activity(7, Request({"type": "http"}))  # idempotent
            app.undelete_activity(7, Request({"type": "http"}))
            self.assertEqual(activity["deleted"], 0)
        statements = [call.args[0] for call in self.sql.execute.call_args_list]
        self.assertEqual(sum(s.startswith("INSERT INTO history") for s in statements), 2)
        self.assertFalse(any(s.startswith("DELETE FROM") for s in statements))
        self.assertEqual(sum(s == "COMMIT" for s in statements), 3)
        self.assertTrue(all(call.kwargs.get("lock") for call in tracker.call_args_list))

    def test_trash_restoration_over_capacity_rolls_back(self):
        def execute(sql, params=()):
            result = MagicMock()
            result.fetchone.return_value = dict(id=7, tracker_id=1, details="Old", deleted=1)
            if sql.startswith("SELECT p.iso_week"):
                result.__iter__.return_value = iter([dict(iso_week=41, month_key="2026-10")])
            return result
        self.sql.execute.side_effect = execute
        with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(app, "get_tracker"):
            with self.assertRaises(HTTPException) as error:
                app.undelete_activity(7, Request({"type": "http"}))
        self.assertEqual(error.exception.status_code, 422)
        statements = [call.args[0] for call in self.sql.execute.call_args_list]
        self.assertIn("ROLLBACK", statements)
        self.assertFalse(any(s.startswith("UPDATE") for s in statements))

    def test_other_employee_cannot_delete_or_restore(self):
        def execute(sql, params=()):
            result = MagicMock()
            result.fetchone.return_value = dict(id=7, tracker_id=1, owner_id=2, deleted=0)
            return result
        self.sql.execute.side_effect = execute
        with patch.object(app, "cur_user", return_value={"id": 1, "role": "employee"}):
            for action in (app.delete_activity, app.undelete_activity):
                with self.assertRaises(HTTPException) as error: action(7, Request({"type": "http"}))
                self.assertEqual(error.exception.status_code, 404)
        self.assertFalse(any(call.args[0].startswith("UPDATE") for call in self.sql.execute.call_args_list))

    def test_deleted_work_is_excluded_from_tracker_and_excel_totals(self):
        import io
        from datetime import date
        from openpyxl import load_workbook
        from export_xlsx import build
        period = dict(id=1, start="2026-10-05", end="2026-10-09", month_key="2026-10", iso_year=2026, iso_week=41, capacity=5)
        active = dict(id=7, tracker_id=1, details="Kept archived", archived=1, status="Completed",
                      affected="", deliverables="", estimation="", progress=60, pos=1, deleted=0)
        deleted = dict(id=8, details="Removed", archived=0, total=1500)
        def execute(sql, params=()):
            result = MagicMock()
            rows = []
            if sql.startswith("SELECT " + app.PSEL): rows = [period]
            elif sql.startswith("SELECT * FROM activities WHERE tracker_id"):
                rows = [active] + ([] if "deleted=0" in sql else [deleted])
            elif sql.startswith("SELECT a.id,a.details"): rows = [deleted]
            elif sql.startswith("SELECT e.*"):
                rows = [dict(activity_id=7, period_id=1, milli=500)]
                if "a.deleted=0" not in sql: rows.append(dict(activity_id=8, period_id=1, milli=1500))
            elif sql.startswith("SELECT MIN"):
                result.fetchone.return_value = dict(first=period["start"], last=period["end"], total=500 if "a.deleted=0" in sql else 2000)
            elif sql.startswith("SELECT id,name,email"): result.fetchone.return_value = dict(id=1, name="Demo", email="demo@example.org")
            result.__iter__.return_value = iter(rows)
            return result
        self.sql.execute.side_effect = execute
        tracker = dict(id=1, owner_id=1, work_package="", start_date="2026-01-01", end_date=None)
        with patch.object(app, "resolve_range", return_value=(date(2026,10,1), date(2026,10,31))), patch.object(app, "affiliation", return_value={}):
            data = app.tracker_data(self.sql, tracker)
        self.assertEqual(data["col_totals"], {1: 500})
        self.assertEqual(data["month_totals"], {"2026-10": 500})
        self.assertEqual(data["bounds"]["all_total"], 500)
        self.assertEqual([a["id"] for a in data["activities"]], [7])
        self.assertEqual(data["deleted_activities"], [deleted])
        wb = load_workbook(io.BytesIO(build({**data, "comments": []})))
        values = [c.value for row in wb["Weekly Status"] for c in row]
        self.assertIn("Kept archived (archived)", values)
        self.assertNotIn("Removed", values)
        self.assertEqual(wb["Monthly BL"].cell(3, 3).value, 0.5)

    def test_stale_save_cannot_edit_deleted_activity(self):
        for kind in ("cell", "field"):
            def execute(sql, params=()):
                result = MagicMock()
                result.fetchone.return_value = None
                if sql.startswith("SELECT " + app.PSEL):
                    result.__iter__.return_value = iter([dict(id=1)])
                return result
            self.sql.execute.side_effect = execute
            with patch.object(app, "cur_user", return_value={"id": 1}), patch.object(app, "get_tracker"):
                with self.assertRaises(HTTPException) as error:
                    app.save(1, app.Save(changes=[dict(type=kind, activity_id=7, period_id=1, field="details", new="changed")]), Request({"type":"http"}))
            self.assertEqual(error.exception.status_code, 404)
        queries = [call.args[0] for call in self.sql.execute.call_args_list if call.args[0].startswith("SELECT * FROM activities")]
        self.assertTrue(all("deleted=0" in sql for sql in queries))

if __name__ == "__main__":
    unittest.main()
